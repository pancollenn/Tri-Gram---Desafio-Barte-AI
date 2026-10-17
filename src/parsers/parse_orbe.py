"""Parser da linha de negócio CHECKOUT (Adquirência Orbe e Taxas Contratadas).

Responsabilidade:
    Ler o relatório de comissões da adquirente Orbe e cruzar com a tabela de taxas
    contratadas vigentes, entregando estruturas de dados normalizadas e tipadas (Decimal).

Entradas:
    - 01_orbe_relatorio_comissoes_*.csv (delimitador ';', vírgula decimal)
    - 02_tabela_taxas_contratadas*.xlsx (aba 'Taxas vigentes', ignorando aba antiga)
    - 06_orbe_ajustes_tardios_*.csv (opcional, delimitador ';', ajustes tardios de fechamento)
"""
from __future__ import annotations

import csv
import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union

import openpyxl

from src import config


def _strip_accents(texto: str) -> str:
    """Remove acentos e converte para minúsculas para comparação textual tolerante."""
    return "".join(
        c for c in unicodedata.normalize("NFD", texto)
        if unicodedata.category(c) != "Mn"
    ).lower().strip()


def _parse_vigencia(texto: Optional[Union[str, date, datetime]]) -> Tuple[Optional[date], Optional[date]]:
    """Extrai intervalo [data_inicio, data_fim] da coluna Vigência.

    Exemplos:
        - 'até 14/09/2026' -> (None, 2026-09-14)
        - 'a partir de 15/09/2026' -> (2026-09-15, None)
        - '01/01/2026' -> (2026-01-01, None)
        - None ou '' -> (None, None)
    """
    if not texto:
        return (None, None)
    if isinstance(texto, datetime):
        return (texto.date(), None)
    if isinstance(texto, date):
        return (texto, None)

    limpo = str(texto).strip()
    m_ate = re.search(r"at[eé]\s*(\d{2}/\d{2}/\d{4})", limpo, re.IGNORECASE)
    if m_ate:
        d = datetime.strptime(m_ate.group(1), "%d/%m/%Y").date()
        return (None, d)

    m_a_partir = re.search(r"a\s+partir\s+de\s*(\d{2}/\d{2}/\d{4})", limpo, re.IGNORECASE)
    if m_a_partir:
        d = datetime.strptime(m_a_partir.group(1), "%d/%m/%Y").date()
        return (d, None)

    m_data = re.search(r"(\d{2}/\d{2}/\d{4})", limpo)
    if m_data:
        d = datetime.strptime(m_data.group(1), "%d/%m/%Y").date()
        return (d, None)

    return (None, None)


def _parse_taxa_celula(val: Any) -> Optional[Decimal]:
    """Converte percentual de célula do Excel (float, str com '%' ou 'padrão') para Decimal.

    Exemplos:
        - 0.0149 (float) -> Decimal('1.49')
        - '1,49%' (str) -> Decimal('1.49')
        - 'padrão' -> config.SPREAD_PADRAO_PCT (Decimal('1.50'))
        - None / '' -> None
    """
    if val is None or val == "":
        return None
    if isinstance(val, (int, float)):
        dec = Decimal(str(val))
        return dec * Decimal("100")
    texto = str(val).strip()
    if texto.lower() in ("padrão", "padrao"):
        return config.SPREAD_PADRAO_PCT
    texto = texto.replace("%", "").strip()
    return config.parse_decimal_ptbr(texto)


def carregar_taxas_contratadas(caminho_taxas: Path) -> List[Dict[str, Any]]:
    """Lê a aba 'Taxas vigentes' do Excel e devolve a lista de contratos estruturada."""
    wb = openpyxl.load_workbook(caminho_taxas, data_only=True)
    nome_aba = "Taxas vigentes"
    if nome_aba not in wb.sheetnames:
        for s in wb.sheetnames:
            if "vigente" in s.lower():
                nome_aba = s
                break
    ws = wb[nome_aba]

    contratos: List[Dict[str, Any]] = []
    # Linha 6 inicia os dados (linhas 1-3 metadados, 4 subcabeçalho, 5 cabeçalho)
    for row in ws.iter_rows(min_row=6, values_only=True):
        lojista = row[0]
        if not lojista or "*" in str(lojista):
            continue
        cnpj_raw = row[1]
        ec_raw = row[2]
        cnpj = config.clean_cnpj(cnpj_raw) if cnpj_raw else None
        ec = str(ec_raw).strip() if ec_raw is not None else None

        deb = _parse_taxa_celula(row[3])
        cred_v = _parse_taxa_celula(row[4])
        cred_2_6 = _parse_taxa_celula(row[5])
        cred_7_12 = _parse_taxa_celula(row[6])
        pix = _parse_taxa_celula(row[7])
        antec = _parse_taxa_celula(row[8])
        spread = _parse_taxa_celula(row[9])

        dt_ini, dt_fim = _parse_vigencia(row[10])
        obs = str(row[11]).strip() if len(row) > 11 and row[11] is not None else ""

        contratos.append({
            "lojista": str(lojista).strip(),
            "cnpj": cnpj,
            "ec": ec,
            "taxas": {
                "debito": deb,
                "credito_a_vista": cred_v,
                "credito_2_6x": cred_2_6,
                "credito_7_12x": cred_7_12,
                "pix": pix,
                "antecipacao": antec,
                "spread": spread,
            },
            "vigencia_inicio": dt_ini,
            "vigencia_fim": dt_fim,
            "obs": obs,
        })
    return contratos


def _obter_taxa_contratada(
    ec: str,
    tipo_lancamento: str,
    modalidade: str,
    qtd_parcelas: Optional[int],
    dt_efetiva: date,
    contratos: List[Dict[str, Any]],
) -> Tuple[str, Optional[str], Optional[Decimal], bool, Optional[str]]:
    """Encontra o lojista, CNPJ e taxa contratada aplicável à transação.

    Returns:
        (lojista, cnpj, taxa_contratada_pct, sem_contrato, modalidade_contrato)
    """
    # 1. EC sem contrato assinado em 2026 (RC-05: Atlas Idiomas)
    if ec in config.EC_SEM_CONTRATO_2026:
        nome_lojista = config.EC_SEM_CONTRATO_2026[ec]
        return nome_lojista, None, None, True, None

    # 2. Casa do Estudante: taxa negociada 0.00% (RC-04)
    if ec == config.EC_CASA_DO_ESTUDANTE:
        candidatos = [c for c in contratos if c["ec"] == ec]
        lojista = candidatos[0]["lojista"] if candidatos else "Casa do Estudante Viagens Ltda"
        cnpj = candidatos[0]["cnpj"] if candidatos else "00520815000100"
        return lojista, cnpj, config.TAXA_CONTRATADA_CASA_DO_ESTUDANTE_PCT, False, "casa_do_estudante"

    # 3. Busca por EC direto
    candidatos = [c for c in contratos if c["ec"] == ec]

    # 4. Fallback: ECs faltantes na tabela mapeados em config (RC-06)
    if not candidatos and ec in config.EC_FALTANTES_TABELA:
        nome_faltante = _strip_accents(config.EC_FALTANTES_TABELA[ec])
        candidatos = [
            c for c in contratos
            if _strip_accents(c["lojista"]) in nome_faltante or nome_faltante in _strip_accents(c["lojista"])
        ]

    if not candidatos:
        return f"EC {ec} (Não mapeado)", None, None, True, None

    # 5. Desempate por vigência temporal (ex.: Bridge Idiomas até 14/09 vs a partir de 15/09)
    contrato = None
    if len(candidatos) == 1:
        contrato = candidatos[0]
    else:
        for c in candidatos:
            ini = c["vigencia_inicio"]
            fim = c["vigencia_fim"]
            if (ini is None or dt_efetiva >= ini) and (fim is None or dt_efetiva <= fim):
                contrato = c
                break
        if not contrato:
            contrato = candidatos[0]

    # 6. Mapeamento da modalidade da transação para a coluna correspondente
    mod_chave = None
    if tipo_lancamento in ("ANTECIPACAO", "REVERSAO ANTECIPACAO"):
        mod_chave = "antecipacao"
    elif modalidade == "DEBITO":
        mod_chave = "debito"
    elif modalidade == "PIX":
        mod_chave = "pix"
    elif modalidade == "CREDITO A VISTA":
        mod_chave = "credito_a_vista"
    elif modalidade == "CREDITO PARCELADO":
        p = qtd_parcelas or 1
        if p <= 1:
            mod_chave = "credito_a_vista"
        elif 2 <= p <= 6:
            mod_chave = "credito_2_6x"
        else:
            mod_chave = "credito_7_12x"

    taxa = contrato["taxas"].get(mod_chave) if mod_chave else None
    return contrato["lojista"], contrato["cnpj"], taxa, False, mod_chave


def _ler_arquivo_transacoes_csv(
    caminho_csv: Path, contratos: List[Dict[str, Any]]
) -> List[Dict[str, Any]]:
    """Lê um arquivo CSV de transações da Orbe e devolve transações normalizadas e enriquecidas."""
    transacoes: List[Dict[str, Any]] = []
    with open(caminho_csv, mode="r", encoding="utf-8-sig") as f:
        leitor = csv.DictReader(f, delimiter=config.SEPARADOR_CSV)
        for r in leitor:
            id_transacao = r["id_transacao"].strip()
            ec = r["cod_estabelecimento"].strip()
            nome_estabelecimento = r["nome_estabelecimento"].strip()
            tipo_lancamento = r["tipo_lancamento"].strip().upper()
            modalidade = r["modalidade"].strip().upper()
            bandeira = r["bandeira"].strip()
            id_referencia = r["id_referencia"].strip()

            # Parcelas
            qtd_parcelas_raw = r.get("qtd_parcelas", "").strip()
            qtd_parcelas = int(qtd_parcelas_raw) if qtd_parcelas_raw else None

            # Valores monetários e percentuais convertidos para Decimal
            valor_bruto = config.parse_decimal_ptbr(r["valor_bruto"])
            taxa_aplicada_pct = config.parse_decimal_ptbr(r["taxa_aplicada_pct"])
            valor_taxa = config.parse_decimal_ptbr(r["valor_taxa"])
            valor_liquido = config.parse_decimal_ptbr(r["valor_liquido"])

            # Datas ISO (AAAA-MM-DD)
            data_movimento = config.parse_data_ptbr(r["data_movimento"])
            dt_mov_obj = datetime.strptime(data_movimento, "%Y-%m-%d").date()

            data_venda_original = ""
            dt_orig_obj: Optional[date] = None
            if r.get("data_venda_original") and r["data_venda_original"].strip():
                data_venda_original = config.parse_data_ptbr(r["data_venda_original"])
                dt_orig_obj = datetime.strptime(data_venda_original, "%Y-%m-%d").date()

            # Data efetiva para determinação da vigência contratual (RC-03: estorno usa venda original)
            dt_efetiva = dt_orig_obj if (tipo_lancamento == "ESTORNO" and dt_orig_obj) else dt_mov_obj
            data_efetiva = dt_efetiva.isoformat()

            lojista, cnpj, taxa_contratada_pct, sem_contrato, mod_contrato = _obter_taxa_contratada(
                ec=ec,
                tipo_lancamento=tipo_lancamento,
                modalidade=modalidade,
                qtd_parcelas=qtd_parcelas,
                dt_efetiva=dt_efetiva,
                contratos=contratos,
            )

            transacoes.append({
                "id_transacao": id_transacao,
                "data_movimento": data_movimento,
                "cod_estabelecimento": ec,
                "nome_estabelecimento": nome_estabelecimento,
                "tipo_lancamento": tipo_lancamento,
                "modalidade": modalidade,
                "qtd_parcelas": qtd_parcelas,
                "bandeira": bandeira,
                "valor_bruto": valor_bruto,
                "taxa_aplicada_pct": taxa_aplicada_pct,
                "valor_taxa": valor_taxa,
                "valor_liquido": valor_liquido,
                "id_referencia": id_referencia,
                "data_venda_original": data_venda_original,
                "data_efetiva": data_efetiva,
                # Enriquecimento contratual:
                "cnpj": cnpj,
                "lojista": lojista,
                "taxa_contratada_pct": taxa_contratada_pct,
                "modalidade_contrato": mod_contrato,
                "sem_contrato": sem_contrato,
                # Arquivo físico de origem desta transação (para o RastroTracker):
                "arquivo_origem": caminho_csv.name,
            })
    return transacoes


def parse_orbe_e_taxas(data_dir: Union[str, Path]) -> Dict[str, Any]:
    """Lê e normaliza o relatório da Orbe, a tabela de taxas e ajustes tardios (se houver).

    Args:
        data_dir: Diretório onde estão os arquivos brutos.

    Returns:
        Dicionário com:
            - 'transacoes': Lista de transações da Orbe consolidadas e enriquecidas.
            - 'taxas_contrato': Mapa de taxas contratadas por EC / lojista.
            - 'arquivo_origem': Nome exato do relatório principal de comissões.
            - 'arquivo_ajustes': Nome exato do arquivo de ajustes tardios (se presente).
    """
    caminho_orbe = config.localizar_arquivo(data_dir, "orbe")
    caminho_taxas = config.localizar_arquivo(data_dir, "taxas")
    caminho_ajustes = config.localizar_arquivo_opcional(data_dir, "orbe_ajustes")

    contratos = carregar_taxas_contratadas(caminho_taxas)

    # Mapa de taxas indexado por EC (e nomes faltantes)
    taxas_contrato: Dict[str, Any] = {}
    for c in contratos:
        if c["ec"]:
            taxas_contrato[c["ec"]] = c
    for ec, nome in config.EC_FALTANTES_TABELA.items():
        if ec not in taxas_contrato:
            nome_clean = _strip_accents(nome)
            match = [
                c for c in contratos
                if _strip_accents(c["lojista"]) in nome_clean or nome_clean in _strip_accents(c["lojista"])
            ]
            if match:
                taxas_contrato[ec] = match[0]

    # Ingestão do relatório principal (01_orbe_relatorio_comissoes_*.csv)
    transacoes_principais = _ler_arquivo_transacoes_csv(caminho_orbe, contratos)

    # Dicionário por id_transacao para permitir atualização ou adição ordenada
    transacoes_map: Dict[str, Dict[str, Any]] = {t["id_transacao"]: t for t in transacoes_principais}

    # Ingestão dos ajustes tardios (06_orbe_ajustes_tardios_*.csv), se existir
    if caminho_ajustes:
        transacoes_ajustes = _ler_arquivo_transacoes_csv(caminho_ajustes, contratos)
        for t in transacoes_ajustes:
            transacoes_map[t["id_transacao"]] = t

    transacoes = list(transacoes_map.values())

    retorno: Dict[str, Any] = {
        "transacoes": transacoes,
        "taxas_contrato": taxas_contrato,
        "arquivo_origem": caminho_orbe.name,
        "caminho_orbe": caminho_orbe,
        "caminho_taxas": caminho_taxas,
    }
    if caminho_ajustes:
        retorno["arquivo_ajustes"] = caminho_ajustes.name
        retorno["caminho_ajustes"] = caminho_ajustes

    return retorno

